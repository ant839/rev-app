"""Turn the FNDN program spreadsheet (ENTIRE PROGRAM tab) into a program-builder plan doc.

Usage: python3 fndn_import.py <file.xlsx> <from YYYY-MM-DD> <out.json> [exercise ids json] [existing plan json]

Reads every phase (or deload / testing block), keeps the weeks from the given Monday on, and writes the doc that
private.builder_save_plan(doc) takes (tracks Strength / Engine / Oly, FNDN group roster, new programs as drafts).
Strength C finishers, Engine and Saturday workouts become conditioning blocks (scored: rounds / time / cals / done / text).

Existing plan json:
  {"id": plan uuid, "programs": {"<block>|<track>": program uuid}}            re-save those blocks in place
  {"id": ..., "append": true, "name": ..., "start_date": ..., "blocks": [...], "roster": {...}}
                                                                              add the sheet's blocks after the plan's
                                                                              existing ones (only the new programs are sent)
"""
import json, re, sys, datetime as dt
import openpyxl

YEAR = 2026
PLAN = {}
FNDN_GROUP = '245f3a87-6085-4eb4-9607-dd68628fd48b'
DAYCOLS = {1: 1, 2: 2, 3: 3, 4: 4, 5: 5, 6: 6}      # sheet column -> day number (Mon=1)
OLYCOLS = {7: 2, 8: 4}                                # Oly Tue, Oly Thu
MON = {m: i + 1 for i, m in enumerate('Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec'.split())}
NUM = {'ONE': 1, 'TWO': 2, 'THREE': 3, 'FOUR': 4, 'FIVE': 5, 'SIX': 6, 'SEVEN': 7, 'EIGHT': 8}

# Sheet name -> exercise name in the app library (anything not here is used as written)
EX_MAP = {'Clean and Jerk': 'Clean & Jerk', 'Conventional Deadlift': 'Deadlift'}
PCT_OF = {'Hang Snatch (Knee)': 'Snatch', 'Hang Clean (Knee)': 'Clean', 'Jerk from Rack': 'Clean & Jerk'}
TRACKS = [('strength', 'Strength', 'strength'), ('engine', 'Engine', 'other'), ('oly', 'Oly', 'technical')]
COND_TITLE = re.compile(r'amrap|emom|tabata|for time|rounds?\b|every \d+ min|plank|core|death by|chipper|interval', re.I)


def cell(ws, r, c):
    v = ws.cell(r, c).value
    return '' if v is None else str(v).strip()


def parse_date(s):
    m = re.match(r'(\d{1,2}) (\w{3})', s.strip())
    return dt.date(YEAR, MON[m.group(2)], int(m.group(1))) if m else None


def clean_note(s):
    s = re.sub(r'\s*—\s*', ', ', s.strip()).strip(' ,')
    return s[:1].upper() + s[1:] if s else s


# ---------- Strength ----------
def strength_a(detail):
    # 1. 'Build to a 3RM — 8 Timed Sets / Every 2 min: Bar x8, 50% x5 ... / Every 3 min: 75% x3, ... Build from here x3'
    m2, m3 = re.search(r'Every 2 min:\s*(.+)', detail), re.search(r'Every 3 min:\s*(.+)', detail)
    if m2 and m3:
        sets, warm = [], []
        for p in [x.strip() for x in m2.group(1).split(',')]:
            mm = re.match(r'(Bar|\d+(?:\.\d+)?%)\s*x(\d+)', p)
            if mm:
                pct = None if mm.group(1) == 'Bar' else float(mm.group(1).rstrip('%'))
                warm.append({'sets': 1, 'reps': mm.group(2), 'percent': pct, 'load_text': 'Bar' if pct is None else None, 'warmup': True})
        for p in [x.strip() for x in m3.group(1).split(',')]:
            mm = re.match(r'(\d+(?:\.\d+)?)%\s*x(\d+)', p)
            if mm:
                sets.append({'sets': 1, 'reps': mm.group(2), 'percent': float(mm.group(1))})
            else:
                mm = re.match(r'(.+?)\s*x(\d+)\s*(\((.+)\))?', p)
                sets.append({'sets': 1, 'reps': mm.group(2), 'load_text': mm.group(1).strip(), 'notes': mm.group(4)})
        sets[0]['notes'] = 'Every 3 min from here'
        return warm + sets, 'Build to a 3RM, 8 timed sets. Build-up sets every 2 min, then every 3 min. The last set is your 3RM attempt.'
    # 2. 'Warmup: Bar / 30% / 40% / 50% — every 2 min (4 sets) / Working: Every 3 min for 3 sets / 3 x 3 @ 60% — note'
    mw = re.search(r'Warmup:\s*Bar\s*((?:/\s*\d+(?:\.\d+)?%\s*)+)', detail)
    mk = re.search(r'(\d+)\s*x\s*(\d+)\s*@\s*(\d+(?:\.\d+)?)%\s*(?:—\s*(.+))?', detail)
    if mw and mk:
        reps = mk.group(2)
        warm = [{'sets': 1, 'reps': reps, 'load_text': 'Bar', 'warmup': True}] + \
               [{'sets': 1, 'reps': reps, 'percent': float(x), 'warmup': True} for x in re.findall(r'(\d+(?:\.\d+)?)%', mw.group(1))]
        ev = re.search(r'Working:\s*Every (\d+ min)', detail)
        note = '. '.join(x for x in [('Every ' + ev.group(1)) if ev else '', clean_note(mk.group(4) or '')] if x)
        return warm + [{'sets': int(mk.group(1)), 'reps': reps, 'percent': float(mk.group(3)), 'notes': note or None}], 'Build-up sets every 2 min.'
    # 3. 'BUILD UP TO A HEAVY SINGLE! / Take your time ...'
    if re.match(r'build (up )?to a heavy single', detail, re.I):
        rest = detail.split('\n', 1)[1].strip() if '\n' in detail else ''
        return [{'sets': 1, 'reps': '1', 'load_text': 'Build to a heavy 1 rep for the day.'}], ('Testing. ' + rest).strip()
    return None


def strength_b(detail):
    m = re.search(r'Working:\s*(\d+)\s*x\s*(\d+)\s*@\s*RPE\s*(\d+(?:\.\d+)?)\s*(?:—\s*every\s*([\d:]+))?', detail)
    if not m:
        return None
    note = ('Every ' + m.group(4) + ', after 3 building sets every 2 min') if m.group(4) else '3 building sets first'
    sub = re.search(r'If reps not achievable:\s*(.+)', detail)
    if sub:
        note += '. If reps not achievable: ' + sub.group(1).strip()
    return [{'sets': int(m.group(1)), 'reps': m.group(2), 'rpe': float(m.group(3)), 'notes': note}], None


def pullup_test(detail, ex_ids):
    """'Pull-Up Progressions: Max added weight / Rest 5 min / Max Bodyweight Reps' -> weighted single + max reps."""
    sub = re.search(r'If bodyweight reps not achievable:\s*(.+)', detail)
    return [item('B1', 'Weighted Pull-Up', ([{'sets': 1, 'reps': '1', 'load_text': 'Build to a heavy 1 rep for the day.'}], 'Testing: max added weight. Enter the added weight, not your bodyweight.'), ex_ids),
            item('B2', 'Pull-Ups (Max Reps)', ([{'sets': 1, 'reps': 'Max', 'load_text': 'Max bodyweight reps',
                                                   'notes': 'Rest 5 min after the weighted single. Enter 0 kg.'}],
                                                 'Testing: max bodyweight reps.' + (' If not achievable: ' + sub.group(1).strip() if sub else '')), ex_ids)]


# ---------- Oly ----------
def oly(detail):
    m = re.match(r'Build to an? (heavy|moderate|light and snappy) (\d+(?:\+\d+)?)', detail)
    if m:
        r = m.group(2)
        word = r + (' rep' if r == '1' else ' reps') if r.isdigit() else r
        return [{'sets': 1, 'reps': r, 'load_text': 'Build to a %s %s for the day.' % (m.group(1), word)}], None
    if re.match(r'Build to a heavy single', detail, re.I):
        rest = re.sub(r'^Build to a heavy single\s*(\(([^)]*)\))?\.?\s*', lambda x: (x.group(2) or '').capitalize() + '. ' if x.group(2) else '', detail, flags=re.I)
        return [{'sets': 1, 'reps': '1', 'load_text': 'Build to a heavy 1 rep for the day.'}], rest.strip() or None
    tail = lambda s: s.split('. ', 1)[1].strip() if '. ' in s else ''
    # 'Light 2s: every 90 sec for 6 sets, 2 reps each at about 60-65% of snatch. Fast and crisp ...'
    m = re.search(r'every (\d+ (?:sec|min)) for (\d+) sets, (?:(\d+) reps|(\d+) clean \+ (\d+) jerk) each at about (\d+)(?:-(\d+))?%', detail)
    if m:
        reps = m.group(3) or '%s+%s' % (m.group(4), m.group(5))
        return [{'sets': int(m.group(2)), 'reps': reps, 'percent': float(m.group(6)), 'percent_max': float(m.group(7)) if m.group(7) else None,
                 'notes': ('Every ' + m.group(1) + '. ' + tail(detail)).strip()}], None
    # 'Light singles @ 75% of current best: 5 singles, every 90 sec. Fast and crisp ...'
    m = re.search(r'singles @ (\d+)% of current best: (\d+) singles, every (\d+ (?:sec|min))', detail)
    if m:
        return [{'sets': int(m.group(2)), 'reps': '1', 'percent': float(m.group(1)), 'notes': ('Every ' + m.group(3) + '. ' + tail(detail)).strip()}], None
    return None


# ---------- Conditioning ----------
def score_for(title, body):
    """How athletes log a conditioning piece: rounds / time / cals / done / text."""
    t, b = title.lower(), body.lower()
    if 'death by' in t: return 'rounds'
    if 'tabata' in t or 'total calories' in b: return 'cals'
    if 'plank' in t: return 'time'
    if 'for time' in t and not re.search(r'\d+\s*x\s*amrap', t): return 'time'   # 'Teams of 2 — For Time (Cap 40)': time, or reps at the cap
    if re.search(r'\d+\s*x\s*amrap', t) or 'part' in t or 'tri test' in t or 'score:' in b or 'split' in b: return 'text'
    if 'amrap' in t: return 'rounds'
    if 'for time' in t or 'for time' in b[:40] or re.match(r'^\d+(-\d+)+', t): return 'time'
    if 'emom' in t or re.match(r'(core\s*—\s*)?every \d+ min', t): return 'done'
    if ('cal' in b or 'bike' in b) and 'max' in b and 'odd' not in b: return 'cals'
    return 'text'


def cond_item(label, title, body, ex_ids):
    """A conditioning block: the Conditioning exercise, the workout in notes, guide.cond = {title, score}."""
    return {'label': label, 'exercise_id': ex_ids.get('conditioning'), 'new_exercise': '' if ex_ids.get('conditioning') else 'Conditioning',
            'percent_of': None, 'notes': body.strip(), 'hold': True,
            'guide': {'off': True, 'cond': {'title': title.strip(), 'score': score_for(title, body)}},
            'sets': [{'sets': 1, 'reps': '1', 'percent': None, 'percent_max': None, 'rpe': None, 'load_text': title.strip(), 'notes': None, 'warmup': False}]}


def item(label, name, rx, ex_ids):
    name = EX_MAP.get(name, name)
    sets, notes = rx
    pof = PCT_OF.get(name)
    it = {'label': label, 'exercise_id': ex_ids.get(name.lower()), 'new_exercise': '' if ex_ids.get(name.lower()) else name,
          'percent_of': ex_ids.get(pof.lower()) if pof else None, 'notes': notes or '', 'hold': True, 'guide': None, 'sets': []}
    for s in sets:
        it['sets'].append({'sets': s.get('sets', 1), 'reps': str(s.get('reps', '1')), 'percent': s.get('percent'), 'percent_max': s.get('percent_max'),
                           'rpe': s.get('rpe'), 'load_text': s.get('load_text'), 'notes': s.get('notes'), 'warmup': bool(s.get('warmup'))})
    return it


# ---------- Sheet ----------
def read_phases(path):
    ws = openpyxl.load_workbook(path, data_only=True)['ENTIRE PROGRAM']
    phases, cur = [], None
    r, R = 1, ws.max_row
    while r <= R:
        a = cell(ws, r, 1)
        # 'PHASE THREE — BUILD TO A 3RM BLOCK  (21 Sep – 17 Oct 2026)' or 'DELOAD & TESTING WEEKS  (19 – 31 Oct 2026)'
        m = re.match(r'^(?:PHASE (\w+)\s*—\s*)?(.+?)\s*\((\d{1,2})(?: (\w{3}))? – (\d{1,2}) (\w{3}) (\d{4})\)\s*$', a)
        if m and not a.startswith('▼') and 'WARMUPS' not in a:
            t = re.sub(r'\s+BLOCK$', '', m.group(2).strip()).lower()
            t = re.sub(r'(\d)rm', r'\1RM', t).replace(' weeks', '')
            start = dt.date(int(m.group(7)), MON[m.group(4) or m.group(6)], int(m.group(3)))
            t = t.title().replace(' And ', ' & ') if not m.group(1) else t[:1].upper() + t[1:]
            cur = {'title': t, 'num': str(NUM.get(m.group(1), m.group(1))) if m.group(1) else None,
                   'start': start, 'weeks': [], 'warm': {}}
            phases.append(cur)
        elif a == 'WARMUP' and cur:  # same warm-up every week of the phase
            for c, d in DAYCOLS.items():
                v = cell(ws, r, c + 1)
                if v and v != '—':
                    cur['warm'][d] = v
        elif cur and re.match(r'^(WEEK \d+|[A-Z][A-Z &]* WEEK)\b', a):
            n = len(cur['weeks']) + 1
            label = re.sub(r'\s+', ' ', a.replace('·', '-')).strip()
            daterow = r + 2
            week = {'n': n, 'label': label, 'days': {}, 'oly': {2: [], 4: []}}
            for c in range(2, 8):
                head = [x.strip() for x in cell(ws, daterow, c).split('|')]
                week['days'][c - 1] = {'date': parse_date(head[0]) if head[0] else None, 'type': head[1] if len(head) > 1 else '',
                                       'code': head[2] if len(head) > 2 else '', 'parts': []}
            rr = daterow + 1
            while rr <= R and cell(ws, rr, 1) in ('A', 'B', 'C', 'D'):
                lab = cell(ws, rr, 1)
                for c in range(2, 8):
                    t, dtl = cell(ws, rr, c), cell(ws, rr + 1, c)
                    if t and t != '—':
                        week['days'][c - 1]['parts'].append((lab, t, dtl))
                for c, d in OLYCOLS.items():
                    t, dtl = cell(ws, rr, c + 1), cell(ws, rr + 1, c + 1)
                    if t and t != '—':
                        week['oly'][d].append((lab, re.sub(r'^[A-D]\.\s*', '', t), dtl))
                rr += 2
            cur['weeks'].append(week)
            r = rr - 1
        r += 1
    return phases


def block_def(name, weeks):
    return {'name': name, 'weeks': weeks, 'focus': '', 'tracks': [
        {'key': k, 'name': n, 'kind': kind, 'mode': 'progressive', 'auto': False, 'inc': 0, 'warmDefault': False,
         'choice': True, 'select': False, 'perVersion': False, 'is_template': False} for k, n, kind in TRACKS]}


def build(path, start_from, ex_ids):
    phases = read_phases(path)
    keep = [p for p in phases if p['start'] + dt.timedelta(weeks=len(p['weeks'])) > start_from]
    append = bool(PLAN.get('append'))
    old_blocks = PLAN.get('blocks', []) if append else []
    plan_start = dt.date.fromisoformat(PLAN['start_date']) if append else keep[0]['start']
    off_weeks = sum(int(b['weeks']) for b in old_blocks)
    if append and plan_start + dt.timedelta(weeks=off_weeks) != keep[0]['start']:
        raise SystemExit('The new block starts %s but the plan\'s existing blocks run to %s' % (keep[0]['start'], plan_start + dt.timedelta(weeks=off_weeks)))
    blocks, programs, problems = list(old_blocks), [], []
    for ph in keep:
        bi = len(blocks)
        bname = ('Phase %s · %s' % (ph['num'], ph['title'])) if ph['num'] else ph['title']
        blocks.append(block_def(bname, len(ph['weeks'])))
        tracks = {'Strength': [], 'Engine': [], 'Oly': []}
        for w in ph['weeks']:
            for d, day in w['days'].items():
                if not day['date'] or day['date'] < start_from or not day['parts'] or 'NO SESSION' in day['type'].upper():
                    continue
                warm = ph['warm'].get(d)
                notes = ('WARM-UP\n' + warm) if warm else ''
                if 'STRENGTH' in day['type']:
                    items = []
                    for lab, t, dtl in day['parts']:
                        if lab == 'A':
                            rx = strength_a(dtl)
                        elif re.match(r'pull-?up progressions', t, re.I):
                            items += pullup_test(dtl, ex_ids); continue
                        elif lab != 'A' and (COND_TITLE.search(t) or not strength_b(dtl)):
                            items.append(cond_item(lab, t, dtl, ex_ids)); continue
                        else:
                            rx = strength_b(dtl)
                        if not rx:
                            problems.append('%s %s %s: could not read "%s"' % (day['date'], lab, t, dtl[:60])); continue
                        items.append(item(lab, t, rx, ex_ids))
                    title = ' + '.join(EX_MAP.get(p[1], p[1]) for p in day['parts'] if p[0] == 'A' or not (COND_TITLE.search(p[1])))
                    tracks['Strength'].append({'week': w['n'], 'day': d, 'title': title, 'notes': notes, 'code': day['code'], 'items': items})
                else:
                    many = len(day['parts']) > 1
                    items = [cond_item(lab if many else '', t, dtl, ex_ids) for lab, t, dtl in day['parts']]
                    tracks['Engine'].append({'week': w['n'], 'day': d, 'title': 'Saturday' if d == 6 else 'Engine', 'notes': notes,
                                             'code': day['code'], 'items': items})
            for d, parts in w['oly'].items():
                date = plan_start + dt.timedelta(days=(off_weeks + w['n'] - 1) * 7 + d - 1)
                if date < start_from or not parts:
                    continue
                items = []
                for lab, t, dtl in parts:
                    rx = oly(dtl)
                    if not rx:
                        problems.append('%s Oly %s %s: could not read "%s"' % (date, lab, t, dtl[:60])); continue
                    items.append(item(lab, t, rx, ex_ids))
                tracks['Oly'].append({'week': w['n'], 'day': d, 'title': 'Oly · ' + ' + '.join(EX_MAP.get(p[1], p[1]) for p in parts),
                                      'notes': '', 'code': 'FNDN P%dW%dD%dO' % (bi + 1, w['n'], d), 'items': items})
        pname = PLAN.get('name') or 'FNDN'
        for k, tn, kind in TRACKS:
            programs.append({'id': PLAN.get('programs', {}).get('%d|%s' % (bi, tn)), 'block_index': bi, 'name': '%s · %s · %s' % (pname, bname, tn),
                             'kind': kind, 'track': tn, 'phase': bname, 'days_version': None, 'is_select': False, 'is_option': True,
                             'is_template': False, 'sessions': tracks[tn]})
        off_weeks += len(ph['weeks'])
    doc = {'id': PLAN.get('id'), 'name': PLAN.get('name') or 'FNDN', 'plan_type': 'FNDN', 'athlete_id': None, 'start_date': plan_start.isoformat(),
           'total_weeks': sum(int(b['weeks']) for b in blocks), 'training_days': PLAN.get('training_days') or [1, 2, 3, 4, 5, 6], 'versions': [],
           'roster': PLAN.get('roster') or {'groups': [FNDN_GROUP], 'members': []}, 'goals': PLAN.get('goals'),
           'code_prefix': PLAN.get('code_prefix') or 'FNDN', 'blocks': blocks, 'programs': programs}
    return doc, problems


if __name__ == '__main__':
    path, frm, out = sys.argv[1], dt.date.fromisoformat(sys.argv[2]), sys.argv[3]
    ex_ids = json.load(open(sys.argv[4])) if len(sys.argv) > 4 else {}
    if len(sys.argv) > 5: PLAN.update(json.load(open(sys.argv[5])))
    doc, problems = build(path, frm, ex_ids)
    json.dump(doc, open(out, 'w'), ensure_ascii=False)
    for p in doc['programs']:
        print('block', p['block_index'] + 1, p['track'], len(p['sessions']), 'sessions', sum(len(s['items']) for s in p['sessions']), 'items')
    print('PROBLEMS:', problems or 'none')
